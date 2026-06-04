from __future__ import annotations

import os
import re
import uuid
from datetime import datetime
from io import BytesIO
from typing import Dict, List, Tuple

from flask import Flask, render_template, request, send_file, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024


def normalize_text(value) -> str:
    if value is None:
        return ""
    text = str(value)
    trans = str.maketrans("ÁÉÍÓÚáéíóúÜüÑñ", "AEIOUaeiouUuNn")
    return re.sub(r"\s+", " ", text.translate(trans)).strip().lower()


def find_header_and_columns(ws) -> Tuple[int, Dict[str, int]]:
    header_row = None
    columns: Dict[str, int] = {}
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        values = [ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1)]
        normalized = [normalize_text(v) for v in values]
        if any("fecha de radicacion" in cell for cell in normalized):
            header_row = row_idx
            for col_idx, cell in enumerate(normalized, start=1):
                if not cell:
                    continue
                if "nit del prestador" in cell:
                    columns["nit"] = col_idx
                elif "razon social del prestador" in cell:
                    columns["institucion"] = col_idx
                elif "fecha de radicacion" in cell:
                    columns["fecha"] = col_idx
                elif "valor neto de la reclamacion" in cell:
                    columns["valor"] = col_idx
                elif "valor saldo de la reclamacion" in cell:
                    columns["saldo"] = col_idx
                elif "numero de reclamo" in cell:
                    columns["reclamo"] = col_idx
            break
    if header_row is None:
        raise ValueError(f"No se encontró la fila de encabezados en la hoja '{ws.title}'.")
    required = {"institucion", "fecha", "valor", "saldo", "reclamo"}
    missing = required - set(columns)
    if missing:
        raise ValueError(f"Faltan columnas requeridas en la hoja '{ws.title}': {', '.join(sorted(missing))}.")
    return header_row, columns


def parse_nit_display(ws, fallback_name: str) -> str:
    candidates: List[str] = []
    for row_idx in range(1, 5):
        for col_idx in range(1, min(ws.max_column, 8) + 1):
            value = ws.cell(row_idx, col_idx).value
            if isinstance(value, str):
                candidates.append(value)
    candidates.append(fallback_name)

    best_digits = ""
    best_score = (-1, -1)  # (length, contains_nit)

    for text in candidates:
        for pattern in [r"nit[:\s]*([0-9\.\-\s]{8,})", r"(\d[\d\.\-\s]{7,}\d)"]:
            for match in re.finditer(pattern, text, flags=re.I):
                raw = match.group(1).strip()
                digits = re.sub(r"\D", "", raw)
                if 9 <= len(digits) <= 10:
                    score = (len(digits), 1 if "nit" in text.lower() else 0)
                    if score > best_score:
                        best_score = score
                        best_digits = digits

    if not best_digits:
        return ""
    if len(best_digits) == 10:
        return f"{best_digits[:-1]} {best_digits[-1]}"
    return best_digits


def summarize_workbook(file_storage) -> Dict:
    content = BytesIO(file_storage.read())
    file_storage.stream.seek(0)
    workbook = load_workbook(content, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row, cols = find_header_and_columns(sheet)

    institution = None
    dates: List[datetime] = []
    cantidad = 0
    valor = 0.0
    saldo = 0.0

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        reclamo = sheet.cell(row_idx, cols["reclamo"]).value
        fecha = sheet.cell(row_idx, cols["fecha"]).value
        valor_cell = sheet.cell(row_idx, cols["valor"]).value
        saldo_cell = sheet.cell(row_idx, cols["saldo"]).value
        inst_cell = sheet.cell(row_idx, cols["institucion"]).value

        if not any(v not in (None, "") for v in [reclamo, fecha, valor_cell, saldo_cell, inst_cell]):
            continue

        if institution is None and inst_cell not in (None, ""):
            institution = str(inst_cell)
        if reclamo not in (None, ""):
            cantidad += 1
        if isinstance(valor_cell, (int, float)):
            valor += float(valor_cell)
        if isinstance(saldo_cell, (int, float)):
            saldo += float(saldo_cell)
        if isinstance(fecha, datetime):
            dates.append(fecha)

    if institution is None:
        institution = os.path.splitext(file_storage.filename or "archivo")[0]

    return {
        "archivo": file_storage.filename or "archivo.xlsx",
        "institucion": institution,
        "nit": parse_nit_display(sheet, file_storage.filename or "archivo.xlsx"),
        "cantidad": cantidad,
        "valor": valor,
        "saldo": saldo,
        "tiene_2026": any(d.year == 2026 for d in dates),
        "min_fecha": min(dates) if dates else None,
        "max_fecha": max(dates) if dates else None,
    }


def build_summary_workbook(records: List[Dict]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
    header_font = Font(name="Times New Roman", size=11, bold=True)
    body_font = Font(name="Times New Roman", size=11)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 17
    ws.column_dimensions["C"].width = 17
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 32

    total_cantidad = sum(item["cantidad"] for item in records)
    total_valor = sum(item["valor"] for item in records)
    total_saldo = sum(item["saldo"] for item in records)

    ws["C1"] = total_cantidad
    ws["D1"] = total_valor
    ws["E1"] = total_saldo
    for cell_ref in ["C1", "D1", "E1"]:
        ws[cell_ref].font = header_font
        ws[cell_ref].alignment = center
        ws[cell_ref].number_format = "#,##0"

    headers = [
        "INSTITUCIÓN",
        "NIT",
        "CANTIDAD DE\nFACTURAS",
        "VALOR FACTURA",
        "CARTERA ACTIVA",
    ]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(2, col_idx, header)
        cell.font = header_font
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    for row_idx, record in enumerate(records, start=3):
        ws.row_dimensions[row_idx].height = 19
        values = [
            record["institucion"],
            record["nit"],
            record["cantidad"],
            record["valor"],
            record["saldo"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = body_font
            cell.border = border
            if col_idx == 1:
                cell.alignment = left
            elif col_idx == 2:
                cell.alignment = center
                cell.number_format = "@"
            else:
                cell.alignment = right
                cell.number_format = "#,##0"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@app.route("/", methods=["GET", "POST"])
def index():
    context = {
        "records": [],
        "download_url": None,
        "has_2026": False,
        "message": None,
        "errors": [],
        "total_cantidad": 0,
        "total_valor": 0,
        "total_saldo": 0,
    }

    if request.method == "POST":
        files = [f for f in request.files.getlist("files") if f and f.filename]
        if not files:
            context["errors"].append("Debes cargar al menos un archivo Excel .xlsx")
            return render_template("index.html", **context)

        records: List[Dict] = []
        for file_storage in files:
            if not file_storage.filename.lower().endswith(".xlsx"):
                context["errors"].append(f"Archivo no permitido: {file_storage.filename}")
                continue
            try:
                records.append(summarize_workbook(file_storage))
            except Exception as exc:
                context["errors"].append(f"{file_storage.filename}: {exc}")

        if records:
            output = build_summary_workbook(records)
            token = uuid.uuid4().hex
            out_path = os.path.join(OUTPUT_DIR, f"resumen_{token}.xlsx")
            with open(out_path, "wb") as fh:
                fh.write(output.getvalue())

            context.update(
                records=records,
                download_url=url_for("download_file", token=token),
                has_2026=any(item["tiene_2026"] for item in records),
                total_cantidad=sum(item["cantidad"] for item in records),
                total_valor=sum(item["valor"] for item in records),
                total_saldo=sum(item["saldo"] for item in records),
            )
            if context["has_2026"]:
                context["message"] = "Se encontraron archivos con Fecha de radicación en 2026."
            else:
                context["message"] = "No se encontraron fechas de radicación en 2026."

    return render_template("index.html", **context)


@app.route("/download/<token>")
def download_file(token: str):
    path = os.path.join(OUTPUT_DIR, f"resumen_{token}.xlsx")
    return send_file(path, as_attachment=True, download_name="Resumen_Cartera.xlsx")


if __name__ == "__main__":
    app.run(debug=True)
